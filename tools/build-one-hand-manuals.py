from pathlib import Path
import json, string
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.opc.constants import RELATIONSHIP_TYPE as RT

OUT = Path(__file__).resolve().parents[1] / 'downloads'
OUT.mkdir(exist_ok=True)

# All practice material is original. Key locations refer to US QWERTY.
LESSONS = [
 dict(title='Finding F and J', new='fj ', goal='Find two tactile landmarks and type with one hand.',
      locate='F and J are on the middle letter row. Most keyboards have a small raised mark on each. Moving right from F, pass G and H to reach J. The wide Spacebar is below the letters.',
      left='Use your left little finger for F and your left index finger for J. Let your thumb find a comfortable place on the Spacebar. If the little finger is difficult to use, move the hand and use a stronger finger.',
      right='Use your right index finger for F and your right little finger for J. Let your thumb find a comfortable place on the Spacebar. If the little finger is difficult to use, move the hand and use a stronger finger.',
      steps=['Locate both raised marks without pressing. Lift the hand, then find the marks again.','Tap F once and release. Tap J once and release. Confirm each letter by touch, speech, or sight.','Type the short lines below. Each visible gap is one Spacebar press. Press Enter after each complete practice line.'],
      practice=['f j f j','ff jj fj jf','fj jf fj jf'],
      check='Can you find F, J, and Space after taking your hand away? Repeat the shortest line until the letters match.'),
 dict(title='Adding G and H', new='gh', goal='Build a central resting position on F G H J.',
      locate='G is immediately right of F. H is immediately left of J. The four central keys run from left to right as F G H J.',
      left='From left to right, place the left little finger on F, ring finger on G, middle finger on H, and index finger on J.',
      right='From left to right, place the right index finger on F, middle finger on G, ring finger on H, and little finger on J.',
      steps=['Find F and J. Rest the two middle fingers lightly on G and H.','Press one key at a time. Let the other fingers relax; they do not need to stay planted.','Type each line twice. Between lines, take the hand away and rebuild the resting position.'],
      practice=['fg hj gh jh','fghj jhgf','fg hg jh gf'],
      check='Say which finger you use for each central key. If you use a different comfortable pattern, record it and keep it consistent.'),
 dict(title='Reaching D and A', new='da', goal='Move the hand toward the left edge and return to the centre.',
      locate='D is one key left of F. To find A from F on the same row, pass D and S, then reach A. S is a landmark today; you will practise it in the next lesson.',
      left='Move your left hand and forearm left together. Use a comfortable index or middle finger for the distant A; use the little finger for D only if it is comfortable. Return using the F bump.',
      right='Move your right hand left as a unit. Use your index finger for D and A, or use the middle finger if that gives better control. Refind F before returning to F G H J.',
      steps=['Feel the route from F to D and back without pressing other keys.','Find A, press it once, release it, and return to the F landmark.','Type the words below slowly. Move the whole hand whenever a reach feels long.'],
      practice=['fd df fa af','dad add fad','a dad had a fad'],
      check='Can you move left and return without guessing? Do not stretch the little finger while keeping the rest of the hand fixed.'),
 dict(title='Completing the home row letters', new='skl', goal='Add S K L and combine the home row letters.',
      locate='S is between A and D. K is one key right of J. L is two keys right of J. Semicolon is farther right; it is taught later.',
      left='For S, move the hand left and use an index or middle finger. For K and L, move right from J and lead with the index finger. Find J again before settling at the centre.',
      right='Use the index finger after moving left to S. For K and L, move the hand right and use an index or middle finger instead of pulling the little finger sideways.',
      steps=['Locate S from D. Practise D S D, returning to F afterward.','Locate K and L from J. Feel J K L and back, then type those keys.','Copy the word line and phrase. A phrase can stay lowercase until capitals are taught.'],
      practice=['ds sd jk kl lj','ask sad glad flask','a lad has a flag'],
      check='Read the phrase back. Check for a single space between words and no accidental semicolons.'),
 dict(title='Adding E and I', new='ei', goal='Use the top letter row to make more words.',
      locate='E is above D and a little to its left on a typical staggered keyboard. I is above K and a little to its left. Rows are offset; they are not straight vertical columns.',
      left='Move left for E and use the index or middle finger. Move right for I and use the index finger. Let the hand travel; return to F or J when you lose your place.',
      right='Use the index finger for E after moving left. Move the hand right for I and choose the index or middle finger. Do not hold the little finger on J while reaching.',
      steps=['Locate D, then feel the key above and slightly left for E. Confirm the letter.','Locate K, then find I above and slightly left. Confirm it before repeating.','Type the words and phrase. Pause between words to regain the central position.'],
      practice=['de ed ki ik','idea file hide life','a lad has a file'],
      check='Check that every I is the letter i. The number 1 has not been introduced.'),
 dict(title='Adding R and T', new='rt', goal='Use two top row keys near the centre.',
      locate='R is above F and slightly left. T is above G and slightly left. E R T appear consecutively from left to right on the top letter row.',
      left='Try the little finger for R and the ring finger for T from the central position. If either is awkward, shift the hand and use an index or middle finger.',
      right='Try the index finger for R and the middle finger for T. Return those fingers to F and G after each short group.',
      steps=['Feel F to R, then G to T. Practise each movement separately.','Copy the short groups with a light tap and full release.','Read or listen to the entire phrase before typing it in small sections.'],
      practice=['fr rf gt tg','red tree gate first','the red flag is at the gate'],
      check='Notice whether you are reaching too far upward and hitting the number row. Refind a tactile landmark if needed.'),
 dict(title='Adding U and O', new='uo', goal='Reach the upper right letter area with one hand.',
      locate='U is above J and slightly left. O is above L and slightly left. The top row sequence in this area is U I O.',
      left='Use the index finger for U near J. For O, move the whole left hand farther right and use an index or middle finger.',
      right='Try the little finger for U only if the movement is comfortable. For O, reposition the hand and use an index or middle finger. Return to J by touch.',
      steps=['Find J and move up to U. Type U, then J, several times.','Find L and move up to O. Confirm it is the letter o.','Copy the words and phrase. Keep the hand moving freely as the words change sides.'],
      practice=['ju uj lo ol','our guide good door','our guide held the door'],
      check='Separate the letter o from zero. Only letters are used in this lesson.'),
 dict(title='Adding W and Y', new='wy', goal='Reach both sides of the top row without twisting the wrist.',
      locate='W is between Q and E, above S and slightly left. Y is between T and U, above H and slightly left. Q is only a landmark here.',
      left='Move left and use an index or middle finger for W. From the centre, try the middle finger for Y and return to H.',
      right='Move left and use the index finger for W. From the centre, try the ring finger for Y; switch to a stronger finger after moving the hand if needed.',
      steps=['Feel E to W and back. Then feel H to Y and back.','Practise the short groups with an even pause between them.','Type the phrase one word at a time, then repeat it as a complete phrase.'],
      practice=['ew we hy yh','way yellow wish style','we saw the yellow flag'],
      check='Is your forearm moving with your hand? Reposition the keyboard if the far left reach pulls your wrist sideways.'),
 dict(title='Adding Q and P', new='qp', goal='Find the two outer letters of the top row.',
      locate='Q is the first letter on the top letter row, immediately right of Tab. P is right of O. Bracket keys lie farther right and are outside this lesson.',
      left='Move left to Q and right to P with the whole hand. Use an index or middle finger for both distant keys. Return to the F or J bump between groups.',
      right='Lead with the index finger for the far left Q. Move right for P and use a comfortable index or middle finger rather than forcing a long little finger stretch.',
      steps=['Find W, move one key left, and confirm Q. Do not press Tab.','Find O, move one key right, and confirm P.','Type the word groups. In ordinary English, Q is often followed by U, but each is a separate key.'],
      practice=['wq qw op po','quit quiet paper proper','please quit the task'],
      check='Can you reach an outer key and find the centre again without searching randomly?'),
 dict(title='Adding V and N', new='vn', goal='Find two keys on the bottom letter row.',
      locate='V is below F and slightly right. N is below H and slightly right. The Spacebar lies below this row.',
      left='Try the little finger for V and the middle finger for N. Move the hand and use a stronger finger if either motion is difficult.',
      right='Try the index finger for V and the ring finger for N. Let the fingers lift freely instead of keeping them pressed on home keys.',
      steps=['Feel F to V and back. Then feel H to N and back.','Type the short groups, listening or looking for accidental spaces.','Copy the words and phrase. Stop at the end of a group before the hand becomes tired.'],
      practice=['fv vf hn nh','van even night never','leave a note on the desk'],
      check='If Space appears instead of a letter, move a little less far down and confirm the key before resuming.'),
 dict(title='Adding B and M', new='bm', goal='Connect the centre and bottom row in useful words.',
      locate='B is between V and N, below G and slightly right. M is right of N, below J and slightly right.',
      left='Try the ring finger for B and the index finger for M. Use a small hand movement if needed; release each key before the next.',
      right='Try the middle finger for B and the little finger for M. A stronger finger after repositioning is a suitable alternative for M.',
      steps=['Locate G to B and J to M. Practise each pair separately.','Type the word list slowly enough to avoid repeated letters.','Copy the phrase, then repeat it with fewer pauses if it remains comfortable.'],
      practice=['gb bg jm mj','book bring home number','bring my book home'],
      check='Listen for B versus V and M versus N. Practise only the pair that is causing difficulty.'),
 dict(title='Completing the alphabet', new='cxz', goal='Add C X Z and practise all 26 letters.',
      locate='C is below D and slightly right. X is below S and slightly right. Z is below A and slightly right. Z X C appear consecutively along the bottom row.',
      left='Move your left hand left for these keys and use a comfortable index or middle finger. Return to F after each short group.',
      right='Move the right hand left as a unit. Lead with the index finger for Z X C; use the middle finger if it improves control.',
      steps=['Find D to C, S to X, and A to Z one route at a time.','Copy the words, then the phrase. Keep letters lowercase.','For an extra review, say any familiar word and type it. You have now practised every alphabet letter.'],
      practice=['dc cx xz zc','box zip six exact','six boxes can fit'],
      check='Choose three difficult letters and locate each three times. There is no speed requirement for finishing the alphabet.'),
 dict(title='Capitals and periods', new='ABCDEFGHIJKLMNOPQRSTUVWXYZ.', goal='Write complete sentences using one hand.',
      locate='Shift is below the home row at the outside edges. The period key is on the bottom row, after comma and before slash. Learn its position by counting two keys right from M.',
      left='Choose the Shift key that your left hand reaches comfortably. For period, move right and use an index or middle finger. Refind J afterward.',
      right='Choose either Shift key by comfort. Move the right hand toward the period and use an index or middle finger if the little finger reach is awkward.',
      steps=['With Sticky Keys on, press and release Shift once. Then press the desired letter for one capital.','Type the rest of the word in lowercase. At the sentence end, press period without Shift.','Use one space before the next sentence. Avoid Caps Lock if it is assigned as your screen reader modifier.'],
      practice=['Sam has a red book.','I can type a full sentence.','My hand can move and return.'],
      check='Check the first capital, the final period, and the spaces. If every letter is capital, check Caps Lock and any latched Shift state.'),
 dict(title='Numbers in short messages', new='1234$', goal='Add 1 2 3 4 and the dollar sign within ordinary writing.',
      locate='The number row is above the top letter row. Its first four digit keys, from left to right, are 1 2 3 4. On a US keyboard, dollar sign is Shift with 4.',
      left='Move the left hand upward and left as needed. Use an index or middle finger for the number row, then use F to find your way back.',
      right='Move the right hand left and upward. Use the index finger for the first four numbers, with the whole hand travelling instead of stretching.',
      steps=['Find 1 through 4 on the top number row. Say and type each once.','For dollar sign with Sticky Keys, press and release Shift, then press 4. Type the amount after it.','Copy the messages. Numbers are included in the words and sentences you are learning to write.'],
      practice=['I have 2 pens and 3 pads.','A pen is $3.','I paid $4 for 1 pad.'],
      check='A dollar sign is not the letter S. Read each amount back character by character.'),
 dict(title='More numbers in daily writing', new='567890', goal='Complete the number row while writing quantities and amounts.',
      locate='Continue right from 4 to 5 6 7 8 9 0. Zero is the last digit key before the minus key. Use the top number row; a separate numeric keypad is not required.',
      left='Move the left hand along the number row in short groups. Lead with the index or middle finger for the rightmost digits and return by the J bump.',
      right='Move your right hand with each group. Use an index or middle finger for accuracy, then return to J or F before typing letters.',
      steps=['Locate 5 and 6, then 7 and 8, then 9 and 0. Practise one pair at a time.','Combine digits into amounts. Period makes the decimal point; no Shift is needed for a digit.','Copy the sentences and read every number in order.'],
      practice=['I have 5 blue pens and 6 red pens.','The book is $7.50.','We need 8 pads and 90 labels.'],
      check='Compare zero with the letter O and one with lowercase l. Confirm characters instead of relying only on their appearance.'),
 dict(title='Commas questions and apostrophes', new=",?'", goal='Add common punctuation to messages.',
      locate='Comma is immediately right of M. Slash is immediately right of period; Shift with slash makes question mark. Apostrophe is right of semicolon on the home row.',
      left='Move the left hand toward the right edge and use an index or middle finger for punctuation. Use a separate Shift press for question mark when Sticky Keys is enabled.',
      right='Move the right hand to the punctuation area and choose a controlled finger. Do not twist the wrist to keep the hand anchored at F G H J.',
      steps=['Locate comma and apostrophe and press each without Shift.','Make question mark by pressing and releasing Shift, then slash.','Copy the sentences. Read punctuation explicitly using your screen reader, Braille display, or visual review.'],
      practice=["Sam's pen is blue.",'I need pens, pads, and tape.','Can we meet at 2?'],
      check='Keep the apostrophe inside the word. Do not put a space before comma, period, or question mark.'),
 dict(title='Times addresses and symbols', new=';:-/@"()+=' , goal='Type useful symbols in a time, an address, and a short note.',
      locate='Semicolon is right of L; Shift makes colon. Apostrophe with Shift makes quotation mark. Minus is right of 0; equals is right of minus. Shift with equals makes plus.',
      left='Use a movable left hand and an index or middle finger for edge symbols. Keep the same finger choice each time when it is comfortable.',
      right='Use a movable right hand and an index or middle finger for edge symbols. Return to the central landmarks whenever you need to reset.',
      steps=['For US symbols, Shift with 2 makes @, Shift with 9 makes left parenthesis, and Shift with 0 makes right parenthesis.','With Sticky Keys, press and release Shift before the symbol key. A plain slash needs no Shift.','Copy the examples. The example email address is practice text; do not send a message.'],
      practice=['Meet at 2:30; bring 3-4 pens.','Practice date 9/20/2026.','sam@example.com','Sam said "yes" (2 + 2 = 4).'],
      check='Check each symbol separately. These combinations depend on the keyboard layout being US QWERTY.'),
 dict(title='Moving and correcting text', new='', goal='Correct a practice sentence without retyping the whole line.',
      locate='The arrow keys are usually at the lower right. Backspace is usually at the top right of the typing area. Delete, Home, and End vary on laptops. Locate them on your own keyboard first.',
      left='Move the left hand to the editing keys as a unit. Use an index or middle finger, then return to F or J by touch. Dedicated editing keys may reduce the need for Fn combinations.',
      right='Move your right hand to the editing keys and return to the centre afterward. Choose a finger that avoids reaching past another finger.',
      steps=['Type the first line. At the end, press Backspace once to remove the final period, then type it again.','Left and Right Arrow move the insertion point. Backspace removes the character before it; Delete removes the character after it.','To fix the second line, press Home, Right Arrow once, Delete once, then type a. The result should match the third line. Use a practice document only.'],
      practice=['Sam has a blue pen.','Sem has a blue pen.','Sam has a blue pen.'],
      check='Home usually goes to the start of the current line and End to its end in a text editor. Check the insertion point before deleting. Reference 3 covers Windows editing keys.'),
 dict(title='Saving copying and reopening', new='', goal='Save your work and use common shortcuts with one hand.',
      locate='Ctrl and Alt are usually on the bottom row. With Sticky Keys on, press and release Ctrl before the next key. Test shortcuts in your practice file.',
      left='Use the left Ctrl key if it is easier for your left hand. After releasing it, move freely to the letter key; you do not need to span both keys.',
      right='Use whichever Ctrl key your right hand reaches comfortably. Release it before moving to the letter key. A right Ctrl key is helpful but not required.',
      steps=['In Notepad, copy the sentence below. Press Ctrl, release, then S. In the save dialog, use Tab to find the folder and File name controls; choose Documents and type the filename below. Activate Save.','Select the sentence with Ctrl then A. Copy with Ctrl then C. Move to its end, press Enter, then paste with Ctrl then V. Undo the paste with Ctrl then Z if needed.','Save again. In Notepad use Ctrl then O, select your saved file in Documents, and activate Open. Read the sentence to verify that you reopened the right file.'],
      practice=['I saved my typing practice.'],
      check='Save and open dialogs can differ by Windows or app version. Use the control labels, not a memorized number of Tab presses. See references 3 and 4.'),
 dict(title='Writing an independent message', new='', goal='Create, check, correct, and save a useful message with one hand.',
      locate='Use your central landmarks and the letter and symbol routes learned in the previous lessons. Keep a short list of your personal finger adaptations nearby.',
      left='Use only the left hand for this task, including Space, Enter, capitals, numbers, and editing. Move the whole hand when needed and pause to reset.',
      right='Use only the right hand for this task, including Space, Enter, capitals, numbers, and editing. Move the whole hand when needed and pause to reset.',
      steps=['Open your practice document and place the cursor on a blank line. Read the complete message first.','Type the message below, then add an original sentence about a familiar daily task.','Review capitals, punctuation, numbers, and spacing. Correct an error, save the file, and reopen it to confirm the saved result.'],
      practice=['Hello Sam,','Can we meet at 2:30?','I have 4 pens and 2 pads.','The total is $8.50.','Thank you.'],
      check='Completion means that you can produce and save useful text with a repeatable, comfortable method. Record the help you used and the next skill you want to practise.')
]


# Lessons 21 to 50 share the website's authoritative copy material and guidance.
site_source = (Path(__file__).resolve().parents[1] / 'keyboarding-one-hand.js').read_text()
site_lessons = json.loads(site_source.split('  const lessons = ', 1)[1].split(';\n\n  const shiftedKeys', 1)[0])
assert len(site_lessons) == 50
for item in site_lessons[20:]:
    entry = dict(item)
    entry['steps'] = [item['instruction'], 'In a practice document, copy each practice line. Use Enter after each line. Review the result before repeating.']
    if item['title'] in ['Connected sentences', 'Sustained paragraph practice']:
        entry['steps'][1] = 'Copy the sentences into your document using the paragraph breaks described above. Review the result before repeating.'
    LESSONS.append(entry)

SOURCES = [
 ('Enable Ireland and ATandMe','Learn One Handed Typing','https://www.atandme.com/learn-one-handed-typing/','Standard keyboards, central positioning, tactile landmarks, and short practice sessions. Published May 27, 2025.'),
 ('Microsoft Support','Make your mouse keyboard and other input devices easier to use','https://support.microsoft.com/en-us/accessibility/windows/make-your-mouse-keyboard-and-other-input-devices-easier-to-use','Windows Sticky Keys, Filter Keys, and keyboard accessibility settings.'),
 ('Microsoft Support','Keyboard shortcuts in Windows','https://support.microsoft.com/en-us/accessibility/windows/keyboard-shortcuts-in-windows','Text editing, navigation, copy, paste, and undo commands.'),
 ('Microsoft Support','Windows keyboard tips and tricks','https://support.microsoft.com/en-us/windows/hardware/input-devices/windows-keyboard-tips-and-tricks','Common Windows and application shortcuts.'),
 ('NV Access','NVDA User Guide','https://download.nvaccess.org/documentation/userGuide.html','Input Help, typing feedback, and screen reader modifier keys.'),
 ('TypingClub and edclub','One hand typing courses','https://www.typingclub.com/','Links to separate left hand and right hand QWERTY courses.'),
 ('TypingTraining','One Handed Typing Training','https://www.typingtraining.com/one-handed-typing.html','A provider example of standard keyboard instruction for one hand.'),
 ('Matias','Half Keyboard and Half QWERTY Pro','https://matias.ca/halfkeyboard/','An alternative that uses a held Spacebar to access mirrored letters; it requires the appropriate hardware.')
]

def p(doc, text='', style=None):
    return doc.add_paragraph(text, style)

def h(doc,text,level=1):
    return doc.add_heading(text,level)

def page(doc,title):
    para=h(doc,title);para.paragraph_format.page_break_before=True

def bullet(doc,text):
    return p(doc,text,'List Bullet')

def steps(doc,items):
    # Separate numbering instances so every procedure begins at 1.
    nroot=doc.part.numbering_part.element
    abstract=OxmlElement('w:abstractNum'); aid=max([int(x.get(qn('w:abstractNumId'))) for x in nroot.findall(qn('w:abstractNum'))]+[0])+1
    abstract.set(qn('w:abstractNumId'),str(aid))
    lvl=OxmlElement('w:lvl');lvl.set(qn('w:ilvl'),'0')
    for tag,val in [('start','1'),('numFmt','decimal'),('lvlText','%1.'),('lvlJc','left')]:
        e=OxmlElement('w:'+tag);e.set(qn('w:val'),val);lvl.append(e)
    pp=OxmlElement('w:pPr');ind=OxmlElement('w:ind');ind.set(qn('w:left'),'360');ind.set(qn('w:hanging'),'360');pp.append(ind);lvl.append(pp)
    abstract.append(lvl);nroot.append(abstract)
    nid=max([int(x.get(qn('w:numId'))) for x in nroot.findall(qn('w:num'))]+[0])+1
    num=OxmlElement('w:num');num.set(qn('w:numId'),str(nid));ai=OxmlElement('w:abstractNumId');ai.set(qn('w:val'),str(aid));num.append(ai);nroot.append(num)
    for text in items:
        para=p(doc,text);np=OxmlElement('w:numPr');il=OxmlElement('w:ilvl');il.set(qn('w:val'),'0');ni=OxmlElement('w:numId');ni.set(qn('w:val'),str(nid));np.append(il);np.append(ni);para._p.get_or_add_pPr().append(np)

def link(doc,label,url):
    para=p(doc);rel=doc.part.relate_to(url,RT.HYPERLINK,is_external=True);hyper=OxmlElement('w:hyperlink');hyper.set(qn('r:id'),rel)
    run=OxmlElement('w:r');pr=OxmlElement('w:rPr');color=OxmlElement('w:color');color.set(qn('w:val'),'000000');pr.append(color);u=OxmlElement('w:u');u.set(qn('w:val'),'single');pr.append(u);run.append(pr);t=OxmlElement('w:t');t.text=label;run.append(t);hyper.append(run);para._p.append(hyper)
    return para

def bookmark(para,name,bid):
    start=OxmlElement('w:bookmarkStart');start.set(qn('w:id'),str(bid));start.set(qn('w:name'),name)
    end=OxmlElement('w:bookmarkEnd');end.set(qn('w:id'),str(bid));para._p.insert(0,start);para._p.append(end)

def jump(doc,label,name):
    para=p(doc,style='Contents Entry');hyper=OxmlElement('w:hyperlink');hyper.set(qn('w:anchor'),name)
    r=OxmlElement('w:r');t=OxmlElement('w:t');t.text=label;r.append(t);hyper.append(r);para._p.append(hyper)

def table(doc, headers, rows, widths):
    t=doc.add_table(rows=1,cols=len(headers));t.alignment=WD_TABLE_ALIGNMENT.CENTER;t.autofit=False
    for c,w in zip(t.columns,widths):c.width=Inches(w)
    pr=t._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ['top','left','bottom','right','insideH','insideV']:
        el=OxmlElement('w:'+edge);el.set(qn('w:val'),'single');el.set(qn('w:sz'),'5');el.set(qn('w:color'),'D9D9D9');borders.append(el)
    pr.append(borders)
    margins=OxmlElement('w:tblCellMar')
    for edge in ['top','left','bottom','right']:
        el=OxmlElement('w:'+edge);el.set(qn('w:w'),'100');el.set(qn('w:type'),'dxa');margins.append(el)
    pr.append(margins)
    for i,label in enumerate(headers):t.rows[0].cells[i].text=label
    repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
    for vals in rows:
        cells=t.add_row().cells
        for i,val in enumerate(vals):cells[i].text=str(val)
    for ri,row in enumerate(t.rows):
        cant=OxmlElement('w:cantSplit');row._tr.get_or_add_trPr().append(cant)
        for ci,cell in enumerate(row.cells):
            cell.width=Inches(widths[ci]);cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'E7EEF5' if ri==0 else 'FFFFFF');cell._tc.get_or_add_tcPr().append(shade)
            for para in cell.paragraphs:
                para.paragraph_format.space_after=Pt(2);para.paragraph_format.space_before=Pt(2)
                for run in para.runs:run.font.size=Pt(12);run.bold=ri==0
    p(doc)
    return t

def setup_doc(hand):
    doc=Document();s=doc.sections[0];s.page_width=Inches(8.5);s.page_height=Inches(11)
    for border in doc.styles.element.iter(qn('w:pBdr')):
        border.getparent().remove(border)
    s.top_margin=Inches(.65);s.bottom_margin=Inches(.65);s.left_margin=Inches(.75);s.right_margin=Inches(.75);s.footer_distance=Inches(.3)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','List Bullet','List Number']:
        st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
        st.font.size=Pt(14);st.paragraph_format.line_spacing=1.12;st.paragraph_format.space_after=Pt(7)
        if st.element.rPr is not None:
            for co in st.element.rPr.findall(qn('w:color')):
                for attr in ['themeColor','themeTint','themeShade']:
                    co.attrib.pop(qn('w:'+attr),None)
    doc.styles['Title'].font.size=Pt(27);doc.styles['Title'].font.bold=True
    doc.styles['Heading 1'].font.size=Pt(20);doc.styles['Heading 1'].paragraph_format.space_after=Pt(12)
    doc.styles['Heading 2'].font.size=Pt(15);doc.styles['Heading 2'].paragraph_format.space_before=Pt(9);doc.styles['Heading 2'].paragraph_format.space_after=Pt(5)
    doc.styles['Subtitle'].font.size=Pt(16)
    st=doc.styles.add_style('Practice',1);st.base_style=doc.styles['Normal'];st.font.name='DejaVu Sans Mono';st.font.size=Pt(15);st.paragraph_format.space_after=Pt(6)
    st=doc.styles.add_style('Contents Entry',1);st.base_style=doc.styles['Normal'];st.font.size=Pt(12.5);st.paragraph_format.space_after=Pt(3)
    st=doc.styles.add_style('Source',1);st.base_style=doc.styles['Normal'];st.font.size=Pt(11.5);st.paragraph_format.space_after=Pt(4)
    lang=OxmlElement('w:lang');lang.set(qn('w:val'),'en-US');doc.styles['Normal'].element.get_or_add_rPr().append(lang)
    footer=s.footer.paragraphs[0];footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    r=footer.add_run('Page ');r.font.size=Pt(10)
    fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
    doc.core_properties.title=f'Typing with the {hand} Hand Only'
    doc.core_properties.subject='A progressive one hand typing manual for a standard US QWERTY keyboard'
    doc.core_properties.author='Accessible Learning Center'
    doc.core_properties.keywords='one hand typing, accessibility, keyboarding, '+hand.lower()+' hand'
    doc.core_properties.comments=''
    return doc

def build(hand):
    key=hand.lower();doc=setup_doc(hand)
    p(doc,f'Typing with the {hand} Hand Only','Title')
    p(doc,'A beginner manual for a standard US keyboard','Subtitle')
    p(doc,'Accessible Learning Center')
    p(doc,'50 progressive lessons with words, sentences, numbers, shortcuts, and timed practice')
    h(doc,'What you will learn',2)
    p(doc,f'This course is for learners who have function in only their {key} hand. Use only your {key} hand to type across the whole keyboard. This manual teaches a movable hand approach, with F G H J as a central resting position. You will learn to find keys, enter text, correct mistakes, and save your work.')
    p(doc,'The lessons assume a physical US English QWERTY keyboard and Windows 11. The letter practice works in other text editors, but settings and shortcuts may differ. No special keyboard or paid course is required.')
    p(doc,'Select Left hand only or Right hand only in Keyboarding Center Settings to load the matching 50 lessons. Lessons 1 to 20 build foundations; Lessons 21 to 50 develop fluency, symbols, and practical writing.')
    h(doc,'How to use this manual',2)
    bullet(doc,'Begin with setup, then follow the lessons in order. Read or hear the instructions before typing the practice lines.')
    bullet(doc,'Type into a separate practice document. Only the text under Practice is copy material; key names and instructions are not.')
    bullet(doc,'Use short sessions and pause as needed. A lesson may take several sessions. Comfort, control, and useful writing matter more than speed.')
    p(doc,'The lesson order, exercises, finger suggestions, and checkpoints are original teaching material. The reference list identifies sources for the general method and software features. Finger choices may be adapted to the learner; this is not a validated rehabilitation protocol.')
    p(doc,'Prepared September 20 2026','Source')

    page(doc,'Table of Contents')
    p(doc,'Use these links or Word heading navigation to move through the manual. Printed page numbers appear at the bottom of every page.','Contents Entry')
    for label,name in [('Getting ready','ready'),('Using one hand for shortcuts','sticky'),('Your hand position and keyboard map','position')]:jump(doc,label,name)
    for i,l in enumerate(LESSONS,1):jump(doc,f'Lesson {i}  {l["title"]}',f'lesson{i}')
    for label,name in [('Timed practice','timed'),('Skills check and answers','assessment'),('Progress record','progress'),('Troubleshooting and alternatives','help'),('Sources and further practice','sources')]:jump(doc,label,name)

    page(doc,'Getting ready');bookmark(doc.paragraphs[-1],'ready',101)
    h(doc,'Set up your workspace',2)
    p(doc,f'Place the letter area within comfortable reach of your {key} hand. Adjust the keyboard angle so the wrist follows the forearm. Leave room for sideways hand movement. A compact keyboard may help positioning; it does not necessarily have smaller keys. [1]')
    p(doc,'Start with a few comfortable minutes. Stop for pain, tingling, or growing strain. If reach or finger control limits practice, an occupational therapist or assistive technology specialist can help adapt the setup.')
    h(doc,'Find the basic controls',2)
    p(doc,'Find Spacebar below the letters, Enter at the right, Backspace above it, and Tab left of the top letter row. F and J usually have raised marks. Confirm each key by speech, touch, or sight; keyboard shapes vary.')
    h(doc,'Open a practice area',2)
    steps(doc,['Press and release the Windows key to open Start. Search for Notepad, then press Enter. A helper can open it during your first session before you know all the letters.','Choose a blank document or new tab if an old note opens. Locate the text editing area; use Tab and listen for an edit control if needed.','Keep this manual separate from the file where you type. Ask a reader to provide one short practice line at a time, or listen to a line and then switch to the practice document.'])
    h(doc,'Make feedback useful',2)
    p(doc,'Use your usual screen reader or Braille display. Start with character feedback and avoid two voices speaking at once. Enlarge text or change contrast for low vision. Written instructions also support learners who do not use audio.')
    p(doc,'NVDA Input Help is toggled with NVDA+1 and announces keys without carrying out their usual actions. Turn it off before typing. The NVDA key is the configured Insert or Caps Lock key. Have an instructor help with this combination if needed. Windows Sticky Keys does not automatically make every screen reader modifier sequential. [5]')

    page(doc,'Using one hand for shortcuts');bookmark(doc.paragraphs[-1],'sticky',102)
    p(doc,'A modifier changes what another key does. Shift creates capitals and some symbols; Ctrl and Alt help run commands. A plus sign in Ctrl+S means a key combination, not a literal plus character.')
    h(doc,'Enable Sticky Keys in Windows 11',2)
    steps(doc,['Open Start, choose Settings, then Accessibility, then Keyboard. A helper may assist with this first setup.','Use Tab to reach the Sticky keys switch. Press Space to turn it on. Check that the switch says On.','Return to your practice editor. Press and release Shift once, then press F. Confirm a capital F. Press F again and confirm a lowercase f.'])
    p(doc,'Sticky Keys allows supported modifier combinations to be entered one key at a time. To save in an editor, press and release Ctrl, then press S. [2]')
    h(doc,'Use a clear sequence',2)
    p(doc,'For this manual, “Shift then letter” means press Shift, release Shift, press the letter, and release the letter. Do not press Shift twice unless you intend to lock it. Some Sticky Keys options change locking or turn the feature off when two keys are pressed together.')
    p(doc,'Check the Sticky Keys options with your instructor. If a modifier stays active, clear it before continuing; if unsure, turn Sticky Keys off and on in Settings and repeat the capital F test. Do not repeatedly press letters while a shortcut seems stuck.')
    h(doc,'Practise only what you need',2)
    p(doc,'The early lessons use lowercase letters. You can wait until Lesson 13 to practise capitals. Setup examples are demonstrations and do not add new letters to the early copy exercises.')
    p(doc,'Filter Keys is a different feature that changes how brief or repeated presses are handled. It can also make fast taps appear to be ignored. Change it only for a specific need and test the result. [2]')
    p(doc,'On a Mac, Chromebook, tablet, or phone, use that system’s accessibility instructions. The Windows setup steps here do not apply to those devices.')

    page(doc,'Your hand position and keyboard map');bookmark(doc.paragraphs[-1],'position',103)
    p(doc,'Use F G H J as a central reference. For distant keys, move the whole hand and find the centre again afterward. Do not hold all four fingers fixed. Enable Ireland describes general central positioning. [1]')
    rows = [('F','Little finger'),('G','Ring finger'),('H','Middle finger'),('J','Index finger'),('Spacebar','Thumb if comfortable')] if key=='left' else [('F','Index finger'),('G','Middle finger'),('H','Ring finger'),('J','Little finger'),('Spacebar','Thumb if comfortable')]
    table(doc,['Central key',f'{hand} hand starting finger'],rows,[1.7,5.3])
    p(doc,'Each lesson suggests fingers for nearby keys. For far keys, move and use an index or middle finger. If fewer fingers are available, use the routes with the finger or approved typing aid you control. Record your adaptation and use it consistently.')
    h(doc,'Rows from left to right',2)
    p(doc,'Top number row: 1 2 3 4 5 6 7 8 9 0')
    p(doc,'Top letter row: Q W E R T Y U I O P')
    p(doc,'Middle letter row: A S D F G H J K L')
    p(doc,'Bottom letter row: Z X C V B N M')
    p(doc,'Use this map for reference. Practise only the keys introduced in your lesson. The rows are staggered, so “above” and “below” include a small sideways offset.')
    h(doc,'Use the same practice routine',2)
    p(doc,'Find the key, press and release, confirm the result, then return when needed. Copy each line twice, using one space per visible gap and Enter at the end. An optional starting goal is 80 percent accuracy on two tries; adjust it to the learner. The early lessons have no minimum words per minute requirement.')

    for i,l in enumerate(LESSONS,1):
        page(doc,f'Lesson {i} {l["title"]}');bookmark(doc.paragraphs[-1],f'lesson{i}',i)
        p(doc,l['goal'])
        h(doc,'Find the keys',2);p(doc,l['locate'])
        h(doc,f'Use your {key} hand',2);p(doc,l[key])
        h(doc,'Step by step',2);steps(doc,l['steps'])
        h(doc,'Practice',2)
        for line in l['practice']:p(doc,line,'Practice')
        if i==19:p(doc,f'Filename for this lesson: {key}typing.txt')
        h(doc,'Check your work',2);p(doc,l['check'])

    page(doc,'Timed practice');bookmark(doc.paragraphs[-1],'timed',104)
    p(doc,'Use timed practice after Lesson 20, when the full passage is familiar. Timing is optional. Ask someone else to run the timer or use an accessible timer. Keep the same text, settings, and correction rule when comparing attempts.')
    h(doc,'Choose a comfortable duration',2)
    bullet(doc,'1 minute: Establish a baseline. Stop when the timer ends, even in the middle of a sentence.')
    bullet(doc,'3 minutes: Try after shorter practice is comfortable. Include reading and correction time in the score.')
    bullet(doc,'6 minutes: Use only if the duration is comfortable. You may stop earlier. Record the actual time; never turn fatigue into a pass or fail test.')
    h(doc,'Passage to copy',2)
    passage='Today I will write a short note. I can move my hand to each key and return to the centre when I need to. My goal is clear, useful writing. I will check the spaces, capitals, and punctuation before I save my work. Sam needs 4 pens and 2 pads for the meeting at 2:30. The pens cost $3.50. Can we bring the blue folder, too? I will read the note again and correct any mistakes. Each careful practice session helps me learn my own pattern.'
    p(doc,passage,'Practice')
    p(doc,'If you finish early, type one space and begin the passage again. Count letters, digits, spaces, and punctuation in the final text. Do not count Enter or correction keys as characters.')
    h(doc,'Calculate speed',2)
    p(doc,'Gross words per minute equals final typed characters divided by 5, then divided by minutes. For example, 150 characters in 3 minutes gives 10 gross words per minute. Report accuracy separately; different typing programs may use different scoring rules.')

    page(doc,'Skills check and answers');bookmark(doc.paragraphs[-1],'assessment',105)
    h(doc,'Practical check',2)
    steps(doc,['Find F, J, Space, Enter, and Backspace without trial and error in an important document.','Type the message from Lesson 50. Add one sentence of your own.','Make a correction, save the file, and reopen it. Record whether you worked independently or with a prompt.'])
    h(doc,'Knowledge check',2)
    for n,t in enumerate(['What should you do when a key is too far to reach comfortably?','How do you make one capital letter with Sticky Keys on?','What does Backspace remove?','How do you type a dollar sign on a US keyboard?','What should you do if keys are announced but no text appears?','Why is the same timing and scoring method useful?'],1):p(doc,f'{n}. {t}')
    h(doc,'Answers',2)
    p(doc,'1. Move the hand and forearm, reposition the keyboard if needed, and return using a landmark. 2. Press and release Shift, then the letter. 3. The character before the insertion point. 4. Shift then 4 with Sticky Keys. 5. Check focus and whether screen reader Input Help or Keyboard Help is on; return to an editable area. 6. It makes attempts comparable.')
    h(doc,'A simple accuracy measure',2)
    p(doc,'For a fixed copy block, count target characters, including spaces and punctuation, as N. Count substitutions, omitted characters, and extra characters as E. Final text accuracy is 100 times (N minus E) divided by N, with a minimum of zero. Corrected mistakes no longer in the final text are not errors in this measure.')
    p(doc,'For a timed attempt, compare only the target portion reached before time ended; do not mark the unattempted remainder wrong. An instructor can align the text after an insertion or omission. Example: 100 target characters with 4 remaining errors gives 96 percent. Record prompts and corrections separately.')

    page(doc,'Progress record');bookmark(doc.paragraphs[-1],'progress',106)
    p(doc,f'Hand used: {key}. Keyboard layout: US QWERTY. Write or dictate your personal finger choices and any settings that helped.')
    table(doc,['Date','Lesson or task','Accuracy','WPM or untimed','Help or next step'],[['','','','',''] for _ in range(8)],[.9,1.65,1.0,1.35,2.1])
    h(doc,'Review after each session',2)
    bullet(doc,'Which key or movement became easier?')
    bullet(doc,'Where did I lose my position, and which landmark helped?')
    bullet(doc,'What will I repeat next time?')
    bullet(doc,'Was the session comfortable, and should it be shorter?')
    p(doc,'Do not compare one hand speed directly with a previous two hand score as a pass or fail rule. Record useful tasks you can complete, not just the number of words per minute.')
    p(doc,'Completion date and instructor comments:')

    page(doc,'Troubleshooting and alternatives');bookmark(doc.paragraphs[-1],'help',107)
    h(doc,'The hand keeps losing its place',2)
    p(doc,'Stop, release all keys, find F or J, and rebuild the resting position. Practise only one outward and return movement. A small removable tactile marker can help if the existing bump is hard to feel; keep it from interfering with key travel.')
    h(doc,'The letters are repeated or ignored',2)
    p(doc,'Use a light tap and release. Check whether a key is being held or whether Filter Keys is changing recognition. In a browser lesson, make sure the input field has focus and the screen reader is in the mode needed for text entry.')
    h(doc,'The lesson sounds confusing',2)
    p(doc,'Use one speech source at a time. Read the entire target, then type one small group. Use character review for punctuation. Instructions such as “press Control to repeat” are specific to some practice programs; they are not general Windows or screen reader commands.')
    h(doc,'A finger or reach is uncomfortable',2)
    p(doc,'Pause and adjust the setup. Use a stronger finger or a smaller practice group. A consistent adapted method is more useful than forcing a prescribed finger. Get individual input for continuing pain, limited movement, or a need for a typing aid.')
    h(doc,'Other methods may suit some learners',2)
    p(doc,'An alternative layout, word prediction, an on-screen keyboard, or speech input may help with a particular task. These require separate setup and evaluation. A visual on-screen keyboard may not be the best access method for someone who cannot see it.')
    p(doc,'Matias Half-QWERTY uses a held Spacebar to reach mirrored letters on compatible hardware. That behaviour is different from an ordinary keyboard. Do not apply that mapping to the exercises in this manual. [8]')
    p(doc,'Any optional tutor should be tried with the learner’s actual screen reader, Braille display, hearing, vision, and motor access before it becomes a required part of instruction.')

    page(doc,'Sources and further practice');bookmark(doc.paragraphs[-1],'sources',108)
    p(doc,'Sources checked September 20 2026. Links provide background or optional practice. Exercises in this manual are original; they do not reproduce another course. Software screens and external courses may change.','Source')
    for i,(org,title,url,desc) in enumerate(SOURCES,1):
        para=p(doc,f'{i}. {org}  {title}','Source');para.paragraph_format.keep_with_next=True
        para=link(doc,'Open '+title,url);para.style=doc.styles['Source'];para.paragraph_format.keep_with_next=True
        p(doc,desc,'Source')
    link(doc,f'Open edclub {key} hand typing course',f'https://www.edclub.com/library/{key}-hand-typing').style=doc.styles['Source']
    p(doc,'The full online course was not accessibility tested for this manual. You can complete the printed or Word exercises without opening it.','Source')

    path=OUT/f'{hand}_Hand_Only_Typing_Manual.docx';doc.save(path)
    return path

def validate_curriculum():
    allowed=set('\n');errors=[]
    for i,l in enumerate(LESSONS,1):
        allowed.update(l['new'])
        for line in l['practice']:
            extra=set(line)-allowed
            if extra:errors.append({'lesson':i,'line':line,'untaught':''.join(sorted(extra))})
    assert set(string.ascii_lowercase)<=set(''.join(l['new'] for l in LESSONS[:12]))
    if errors:raise ValueError(json.dumps(errors))
    return {'lessons':len(LESSONS),'alphabet_complete_by':12,'untaught_practice_characters':0}

if __name__=='__main__':
    result=validate_curriculum();paths=[build('Left'),build('Right')]
    print(json.dumps({'validation':result,'files':[{'path':str(x),'bytes':x.stat().st_size} for x in paths]},indent=2))
